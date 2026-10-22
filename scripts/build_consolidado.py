# -*- coding: utf-8 -*-
"""Excel consolidado de deportivos y alta gama usados en Colombia."""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILES = ["chevrolet", "ford", "mercedes", "audi", "mini", "porsche", "otros"]

# --- BMW Z4 (recolectado en la fase anterior) ---
z4 = [
 ("BMW","Z4","TuCarro","BMW Z4 sDrive20i",206000000,None,"G29 (2019+) sDrive20i",1998,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-3022543068-bmw-z4-sdrive20i-_JM","Cilindraje derivado de la versión sDrive20i (motor 2.0T B48)"),
 ("BMW","Z4","TuCarro","BMW Z4 M40i Roadster",335000000,None,"G29 (2019+) M40i",2998,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-1261722433-bmw-z4-m40i-roadster-_JM","Cilindraje derivado de la versión M40i (motor 3.0T B58)"),
 ("BMW","Z4","TuCarro","BMW Z4 2.0 sDrive20i",105000000,None,"E89 (2009-2016) sDrive20i",1997,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-2867800666-bmw-z4-20-sdrive20i-_JM","Cilindraje derivado de la versión sDrive20i E89 (motor 2.0T N20)"),
 ("BMW","Z4","TuCarro","BMW Z4 M40i 3.0 2022",290000000,2022,"G29 (2019+) M40i",2998,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-4184920468-bmw-z4-m40i-30-2022-_JM","Cilindraje derivado de la versión M40i (motor 3.0T B58)"),
 ("BMW","Z4","TuCarro","BMW Z4 sDrive20i",None,None,"G29 (2019+) sDrive20i",1998,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-3264022100-bmw-z4-sdrive20i-_JM","Precio no visible en el resultado de búsqueda"),
 ("BMW","Z4","TuCarro","BMW Z4 sDrive20i 2015",120000000,2015,"E89 (2009-2016) sDrive20i",1997,None,None,"No indicado","https://articulo.tucarro.com.co/MCO-1530792721-bmw-z4-sdrive-20i-2015-_JM","Cilindraje derivado de la versión sDrive20i E89 (motor 2.0T N20)"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 sDrive 2.0 Convertible",198900000,None,"G29 (2019+) sDrive20i",1998,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-2864275518-bmw-z4-sdrive-20-convertible-_JM","Cilindraje derivado de la versión sDrive 2.0"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 M40i (placa JQU291)",280000000,None,"G29 (2019+) M40i",2998,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-909293961-bmw-z4-m40i-jqu291-_JM","Cilindraje derivado de la versión M40i (motor 3.0T B58)"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 M40i Convertible 2022",249900000,2022,"G29 (2019+) M40i",2998,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-1528740285-bmw-z4-m40i-convertible-2022-_JM","Cilindraje derivado de la versión M40i (motor 3.0T B58)"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 (versión no especificada)",196000000,None,"ND",None,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-1588628127-bmw-z4-_JM","El título del aviso no especifica versión ni año"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 E86 Coupé 3.0si",198000000,None,"E86 (2006-2008) 3.0si",2996,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-2913422548-bmw-z4-e86-coupe-30si-_JM","Cilindraje derivado de la versión 3.0si (motor N52)"),
 ("BMW","Z4","MercadoLibre CO","BMW Z4 2.5 sDrive23i",107000000,None,"E89 (2009-2016) sDrive23i",2497,None,None,"No indicado","https://carro.mercadolibre.com.co/MCO-2836955948-bmw-z4-_JM","Cilindraje derivado de la versión sDrive23i (motor 2.5 N52)"),
 ("BMW","Z4","CarroYa","BMW Z4 2022",None,2022,"G29 (2019+)",None,None,"Medellín","No indicado","https://www.carroya.com/detalle/usado/bmw/z4/2022/medellin/7a0917d0-534c-42da-ba82-0ca33e8c677e","Portal CarroYa; precio y km se cargan por JavaScript y no fueron accesibles"),
]

rows = []
for m, mo, f, t, p, a, v, cc, km, ciu, rep, link, nota in z4:
    rows.append(dict(marca=m, modelo=mo, fuente=f, titulo=t, precio_cop=p, anio=a,
                     version=v, cilindraje_cc=cc, km=km, ciudad=ciu,
                     para_reparar=rep, link=link, nota=nota))
for name in FILES:
    with open(f"/home/claude/data_{name}.json", encoding="utf-8") as fh:
        rows.extend(json.load(fh))

# --- Deduplicación por ID de publicación (MCO-xxxxx) ---
seen = {}
for r in rows:
    m = re.search(r"MCO-(\d+)", r["link"])
    r["_id"] = m.group(1) if m else None
    if r["_id"] and r["_id"] in seen:
        r["registro"] = "Duplicado (mismo aviso en otro portal)"
    else:
        r["registro"] = "Único"
        if r["_id"]:
            seen[r["_id"]] = True

# Orden: marca, modelo, precio desc
rows.sort(key=lambda r: (r["marca"], r["modelo"], -(r["precio_cop"] or 0)))

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BASE = Font(name="Arial", size=10)
SMALL = Font(name="Arial", size=9)
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="595959")
LINK_FONT = Font(name="Arial", size=9, color="0563C1", underline="single")
DUP_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

wb = Workbook()

# ================= Hoja 1: Avisos =================
ws = wb.active
ws.title = "Avisos"
headers = ["#", "Marca", "Modelo", "Fuente", "Título del aviso", "Precio venta (COP)",
           "Año modelo", "Versión / generación", "Cilindraje (cc)", "Kilometraje (km)",
           "Ciudad", "¿Para reparar?", "Registro", "Link del aviso", "Nota / supuesto",
           "Precio estadístico (COP)"]
ws.append(headers)
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill, cell.font = HDR_FILL, HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN

ND = "ND"
for i, r in enumerate(rows, start=1):
    rr = i + 1
    vals = [i, r["marca"], r["modelo"], r["fuente"], r["titulo"],
            r["precio_cop"] if r["precio_cop"] is not None else ND,
            r["anio"] if r["anio"] is not None else ND,
            r.get("version") or ND,
            r["cilindraje_cc"] if r["cilindraje_cc"] is not None else ND,
            r["km"] if r["km"] is not None else ND,
            r["ciudad"] if r["ciudad"] else ND,
            r["para_reparar"], r["registro"], r["link"], r["nota"],
            f'=IF(AND(M{rr}="Único",ISNUMBER(F{rr})),F{rr},"")']
    ws.append(vals)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=rr, column=c)
        cell.font = BASE
        cell.border = THIN
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 8, 15)))
    ws.cell(row=rr, column=6).number_format = '$#,##0'
    ws.cell(row=rr, column=10).number_format = '#,##0'
    lc = ws.cell(row=rr, column=14)
    lc.hyperlink = r["link"]
    lc.font = LINK_FONT
    lc.alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(row=rr, column=15).font = NOTE_FONT
    ws.cell(row=rr, column=13).font = SMALL
    ws.cell(row=rr, column=16).number_format = '$#,##0'
    if r["registro"] != "Único":
        for c in range(1, len(headers) + 1):
            ws.cell(row=rr, column=c).fill = DUP_FILL

widths = [4, 15, 18, 16, 34, 17, 9, 30, 12, 13, 11, 12, 16, 46, 44, 20]
for c, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:P{len(rows)+1}"

n = len(rows)
last = n + 1
nr = last + 2
ws.cell(row=nr, column=1, value=(
    "ND = dato no disponible públicamente en el resultado de búsqueda. TuCarro, MercadoLibre, CarroYa, Mitula, "
    "VendeTuNave y demás portales colombianos bloquean la lectura automatizada de sus páginas (error 403 / robots.txt), "
    "por lo que kilometraje y ciudad no pudieron extraerse aviso por aviso; los datos provienen de los resultados de "
    "búsqueda indexados de cada publicación. El cilindraje se derivó de la versión indicada en el título de cada aviso "
    "cuando no venía explícito — ver la columna de notas. Las filas resaltadas en amarillo son la misma publicación "
    "espejada entre TuCarro y MercadoLibre (mismo ID MCO) y se excluyen de las estadísticas. "
    "Recolección: 24 de agosto de 2026.")).font = NOTE_FONT
ws.merge_cells(start_row=nr, start_column=1, end_row=nr + 3, end_column=16)
ws.cell(row=nr, column=1).alignment = Alignment(wrap_text=True, vertical="top")

# ================= Hoja 2: Análisis por marca =================
wa = wb.create_sheet("Análisis por marca")
wa.column_dimensions["A"].width = 30
for col in "BCDEFG":
    wa.column_dimensions[col].width = 19

def title_cell(sheet, row, text):
    c = sheet.cell(row=row, column=1, value=text)
    c.font = Font(name="Arial", bold=True, size=11, color="1F4E79")

price = f"Avisos!F2:F{last}"
marca = f"Avisos!B2:B{last}"
modelo = f"Avisos!C2:C{last}"
reg = f"Avisos!M2:M{last}"
pstat = f"Avisos!P2:P{last}"
UNI = '"Único"'

title_cell(wa, 1, "Caracterización de precios — deportivos y alta gama usados en Colombia (24-ago-2026)")

wa.cell(row=3, column=1, value="Avisos totales recolectados").font = BASE
wa.cell(row=3, column=2, value=f"=COUNTA(Avisos!A2:A{last})").number_format = '0'
wa.cell(row=4, column=1, value="Avisos únicos (sin espejos entre portales)").font = BASE
wa.cell(row=4, column=2, value=f'=COUNTIF({reg},{UNI})').number_format = '0'
wa.cell(row=5, column=1, value="Avisos únicos con precio publicado").font = BASE
wa.cell(row=5, column=2, value=f'=COUNTIFS({reg},{UNI},{price},">0")').number_format = '0'
wa.cell(row=6, column=1, value="Avisos anunciados como chocados / para reparar").font = BASE
wa.cell(row=6, column=2, value=f'=COUNTIF(Avisos!L2:L{last},"Sí")').number_format = '0'
for r in range(3, 7):
    wa.cell(row=r, column=1).font = BASE
    wa.cell(row=r, column=2).font = BASE
    for c in (1, 2):
        wa.cell(row=r, column=c).border = THIN

wa.cell(row=7, column=1, value="Mediana general de precios").font = BASE
wa.cell(row=7, column=2, value=f'=MEDIAN({pstat})').number_format = '$#,##0'
wa.cell(row=7, column=2).font = BASE
for c in (1, 2):
    wa.cell(row=7, column=c).border = THIN

title_cell(wa, 9, "Precios por marca (solo avisos únicos con precio)")
mh = ["Marca", "N.º avisos", "Precio mínimo", "Precio promedio", "Precio máximo"]
for c, h in enumerate(mh, start=1):
    cell = wa.cell(row=10, column=c, value=h)
    cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, THIN
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

marcas = sorted({r["marca"] for r in rows})
r0 = 11
for i, mk in enumerate(marcas):
    r = r0 + i
    q = f'"{mk}"'
    wa.cell(row=r, column=1, value=mk)
    wa.cell(row=r, column=2, value=f'=COUNTIFS({marca},{q},{reg},{UNI},{price},">0")')
    wa.cell(row=r, column=3, value=f'=IFERROR(_xlfn.MINIFS({price},{marca},{q},{reg},{UNI}),"—")')
    wa.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIFS({price},{marca},{q},{reg},{UNI}),"—")')
    wa.cell(row=r, column=5, value=f'=IFERROR(_xlfn.MAXIFS({price},{marca},{q},{reg},{UNI}),"—")')
    for c in range(1, 6):
        cell = wa.cell(row=r, column=c)
        cell.font, cell.border = BASE, THIN
        cell.number_format = '$#,##0' if c >= 3 else ('0' if c == 2 else 'General')

end_marcas = r0 + len(marcas) - 1
tot = end_marcas + 1
wa.cell(row=tot, column=1, value="TOTAL / general").font = Font(name="Arial", bold=True, size=10)
wa.cell(row=tot, column=2, value=f'=COUNTIFS({reg},{UNI},{price},">0")')
wa.cell(row=tot, column=3, value=f'=IFERROR(_xlfn.MINIFS({price},{reg},{UNI}),"—")')
wa.cell(row=tot, column=4, value=f'=IFERROR(AVERAGEIFS({price},{reg},{UNI}),"—")')
wa.cell(row=tot, column=5, value=f'=IFERROR(_xlfn.MAXIFS({price},{reg},{UNI}),"—")')
for c in range(1, 6):
    cell = wa.cell(row=tot, column=c)
    cell.font = Font(name="Arial", bold=True, size=10)
    cell.fill, cell.border = SUB_FILL, THIN
    cell.number_format = '$#,##0' if c >= 3 else ('0' if c == 2 else 'General')

# ---- por modelo ----
mstart = tot + 3
title_cell(wa, mstart - 1, "Precios por modelo (solo avisos únicos con precio, ordenado por promedio)")
mh2 = ["Marca y modelo", "N.º avisos", "Precio mínimo", "Precio promedio", "Precio máximo"]
for c, h in enumerate(mh2, start=1):
    cell = wa.cell(row=mstart, column=c, value=h)
    cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, THIN
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

pares = {}
for r in rows:
    if r["registro"] == "Único" and r["precio_cop"]:
        pares.setdefault((r["marca"], r["modelo"]), []).append(r["precio_cop"])
orden = sorted(pares.items(), key=lambda kv: -sum(kv[1]) / len(kv[1]))
for i, ((mk, mo), _) in enumerate(orden):
    r = mstart + 1 + i
    qm, qo = f'"{mk}"', f'"{mo}"'
    wa.cell(row=r, column=1, value=f"{mk} {mo}")
    wa.cell(row=r, column=2, value=f'=COUNTIFS({marca},{qm},{modelo},{qo},{reg},{UNI},{price},">0")')
    wa.cell(row=r, column=3, value=f'=IFERROR(_xlfn.MINIFS({price},{marca},{qm},{modelo},{qo},{reg},{UNI}),"—")')
    wa.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIFS({price},{marca},{qm},{modelo},{qo},{reg},{UNI}),"—")')
    wa.cell(row=r, column=5, value=f'=IFERROR(_xlfn.MAXIFS({price},{marca},{qm},{modelo},{qo},{reg},{UNI}),"—")')
    for c in range(1, 6):
        cell = wa.cell(row=r, column=c)
        cell.font, cell.border = BASE, THIN
        cell.number_format = '$#,##0' if c >= 3 else ('0' if c == 2 else 'General')

# ---- por cilindraje ----
cstart = mstart + len(orden) + 3
title_cell(wa, cstart - 1, "Precios por rango de cilindraje")
for c, h in enumerate(["Rango de cilindraje", "N.º avisos", "Precio mínimo", "Precio promedio", "Precio máximo"], start=1):
    cell = wa.cell(row=cstart, column=c, value=h)
    cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, THIN
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
cc = f"Avisos!I2:I{last}"
rangos = [("Hasta 2.000 cc", 0, 2000), ("2.001 - 3.000 cc", 2001, 3000),
          ("3.001 - 4.000 cc", 3001, 4000), ("4.001 - 5.000 cc", 4001, 5000),
          ("Más de 5.000 cc", 5001, 9999)]
for i, (lbl, lo, hi) in enumerate(rangos):
    r = cstart + 1 + i
    wa.cell(row=r, column=1, value=lbl)
    wa.cell(row=r, column=2, value=f'=COUNTIFS({cc},">="&{lo},{cc},"<="&{hi},{reg},{UNI},{price},">0")')
    wa.cell(row=r, column=3, value=f'=IFERROR(_xlfn.MINIFS({price},{cc},">="&{lo},{cc},"<="&{hi},{reg},{UNI}),"—")')
    wa.cell(row=r, column=4, value=f'=IFERROR(AVERAGEIFS({price},{cc},">="&{lo},{cc},"<="&{hi},{reg},{UNI}),"—")')
    wa.cell(row=r, column=5, value=f'=IFERROR(_xlfn.MAXIFS({price},{cc},">="&{lo},{cc},"<="&{hi},{reg},{UNI}),"—")')
    for c in range(1, 6):
        cell = wa.cell(row=r, column=c)
        cell.font, cell.border = BASE, THIN
        cell.number_format = '$#,##0' if c >= 3 else ('0' if c == 2 else 'General')

# ---- Hallazgos ----
fstart = cstart + len(rangos) + 3
title_cell(wa, fstart - 1, "Hallazgos")
findings = [
 "• Se recolectaron 310 avisos activos de deportivos y alta gama en Colombia (308 únicos tras descontar 2 publicaciones espejadas entre TuCarro y MercadoLibre); 227 publican precio.",
 "• El rango total va de $41,5 millones (MINI Cooper R56 1.6) a $1.150 millones (Porsche 911 Carrera S 992). La mediana general se sitúa en $195 millones.",
 "• Por marca, Porsche encabeza el promedio ($518 M) seguido de Jaguar ($465 M) y Chevrolet ($464 M, empujado por los Corvette C8). En el otro extremo, MINI promedia $103 M y es la marca con más oferta junto a Mercedes-Benz.",
 "• Los modelos más caros del mercado colombiano son Audi R8 V10, Mercedes-AMG GT 53, Corvette (promedio $669 M) y Porsche 911 (promedio $657 M).",
 "• La franja de mayor liquidez está entre $90 y $200 millones: A45/CLA45 AMG, Audi S3 y TT, MINI Cooper S y JCW, BMW Z4 de generaciones anteriores y Mercedes SLK.",
 "• El cilindraje separa claramente el mercado: los motores de hasta 2.000 cc (102 avisos) promedian $137 M, mientras los de más de 5.000 cc (44 avisos) promedian $449 M.",
 "• Ninguno de los 310 avisos se anuncia como chocado o 'para reparar'. Las búsquedas específicas de vehículos siniestrados en Colombia no arrojaron avisos reales de estas marcas.",
 "• Limitación metodológica: los portales colombianos bloquean la lectura automatizada (403 / robots.txt), por lo que kilometraje y ciudad casi nunca están disponibles por aviso, y el cilindraje se derivó de la versión indicada en el título (documentado aviso por aviso en la columna de notas).",
]
for i, f in enumerate(findings):
    r = fstart + i
    cell = wa.cell(row=r, column=1, value=f)
    cell.font = BASE
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    wa.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    wa.row_dimensions[r].height = 32

wb.save("/home/claude/deportivos_alta_gama_colombia.xlsx")
print("filas:", len(rows), "| duplicados:", sum(1 for r in rows if r["registro"] != "Único"),
      "| marcas:", len(marcas), "| modelos:", len(orden))
print("ultima_fila_avisos:", last, "| fila_totales:", tot)
