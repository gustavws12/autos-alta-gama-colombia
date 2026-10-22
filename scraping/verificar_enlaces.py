#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Verificador de enlaces del Excel de avisos.

Para que sirve
--------------
Un aviso de MercadoLibre que ya se vendio, vencio o fue retirado NO devuelve
un error 404: la pagina redirige al listado de carros parecidos y agrega al
final de la URL la marca

    #redirectedFromVip=<la url original del aviso>

Por eso "el link no entra": el navegador si abre algo, pero no es el aviso.
Esta rutina detecta justo eso, sin necesidad de credenciales de la API.

Como decide
-----------
Hay dos modos y la rutina elige el mejor disponible:

MODO API (exacto, si credenciales.json esta lleno)
  Le pregunta a la API el campo `status` de cada aviso. Es la respuesta
  autoritativa: active = vigente, cualquier otra cosa = caido. Ademas es
  rapido: 20 avisos por llamada.

MODO HTTP (aproximado, sin credenciales)
  Pide la URL guardada siguiendo redirecciones y mira DONDE termino:
    - la URL final todavia tiene el codigo MCO del aviso  -> VIGENTE
    - la URL final ya no tiene ese codigo                 -> CAIDO (redirigido)
    - el servidor responde 404 o 410                      -> CAIDO
    - no hubo respuesta (sin internet, bloqueo, timeout)  -> SIN VERIFICAR

  Limitacion honesta del modo HTTP: si MercadoLibre decide mostrar el aviso
  cerrado con un aviso de "publicacion finalizada" en la misma URL, en vez de
  redirigir, este modo lo dara por vigente. El modo API no tiene ese problema.

"Sin verificar" nunca se confunde con "caido": si no hay certeza, se dice.

Uso
---
    python verificar_enlaces.py                    # verifica y marca el Excel
    python verificar_enlaces.py --solo-reporte     # no toca el Excel
    python verificar_enlaces.py --archivo otro.xlsx
    python verificar_enlaces.py --hilos 4 --pausa 0.5

Deja dos columnas nuevas en la hoja "Avisos":
    "Estado del enlace"    Vigente / Caido / Sin verificar
    "Enlace verificado el" fecha de la revision
y un CSV con el detalle en  reporte_enlaces.csv
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import threading
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
XLSX_POR_DEFECTO = RAIZ / "deportivos_alta_gama_colombia.xlsx"
REPORTE = AQUI / "reporte_enlaces.csv"

VIGENTE = "Vigente"
CAIDO = "Caído"
SIN_VERIFICAR = "Sin verificar"

NAVEGADOR = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
             "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")

_candado = threading.Lock()


# ------------------------------------------------------------------ decision
def codigo_mco(url: str) -> str | None:
    """Extrae el codigo del aviso (MCO1542509431) de cualquier URL suya."""
    m = re.search(r"MCO-?(\d{3,})", url or "", re.IGNORECASE)
    return f"MCO{m.group(1)}" if m else None


def sin_fragmento(url: str) -> str:
    """
    Quita todo lo que va despues de '#'.

    Es imprescindible: cuando MercadoLibre redirige un aviso muerto, mete la
    URL ORIGINAL dentro del fragmento

        …/ford/modelo-mustang-shelby/2021/#redirectedFromVip=…MCO-1542509431…

    Si se buscara el codigo en la URL completa se encontraria el del aviso
    muerto y se concluiria que sigue vivo, que es justo lo contrario.
    """
    return (url or "").split("#", 1)[0]


def clasificar(url_original: str, url_final: str | None,
               codigo_http: int | None, error: str | None) -> tuple[str, str]:
    """
    Devuelve (estado, explicacion). Funcion pura: toda la logica de decision
    esta aqui para poder probarla sin red.
    """
    if error:
        return SIN_VERIFICAR, error
    if codigo_http in (404, 410):
        return CAIDO, f"El servidor respondió {codigo_http}"
    if codigo_http is not None and codigo_http >= 400:
        return SIN_VERIFICAR, f"El servidor respondió {codigo_http}"

    codigo = codigo_mco(sin_fragmento(url_original))
    if not codigo:
        return SIN_VERIFICAR, "El enlace guardado no tiene código de aviso"
    if not url_final:
        return SIN_VERIFICAR, "No se pudo leer la dirección final"

    # Se revisa PRIMERO la huella de la redirección: es la señal inequívoca y
    # aparece justo cuando el fragmento todavía contiene el código original.
    if "redirectedfromvip" in url_final.lower():
        return CAIDO, "Redirigido al listado: el aviso ya no existe"

    if codigo_mco(sin_fragmento(url_final)) == codigo:
        return VIGENTE, "El aviso sigue publicado"
    return CAIDO, f"Redirigido a otra página ({_corto(sin_fragmento(url_final))})"


def _corto(url: str, n: int = 70) -> str:
    return url if len(url) <= n else url[:n] + "…"


# --------------------------------------------------------------------- red
def revisar(url: str, tiempo_limite: int = 20) -> tuple[str, str]:
    """Pide la URL siguiendo redirecciones y clasifica donde termino."""
    for metodo in ("HEAD", "GET"):
        pet = urllib.request.Request(url, method=metodo, headers={
            "User-Agent": NAVEGADOR,
            "Accept": "text/html,application/xhtml+xml",
            "Accept-Language": "es-CO,es;q=0.9",
        })
        try:
            with urllib.request.urlopen(pet, timeout=tiempo_limite) as resp:
                return clasificar(url, resp.geturl(), resp.status, None)
        except urllib.error.HTTPError as e:
            # Algunos servidores no aceptan HEAD: se reintenta con GET.
            if metodo == "HEAD" and e.code in (400, 403, 405, 501):
                continue
            destino = getattr(e, "url", None) or url
            return clasificar(url, destino, e.code, None)
        except urllib.error.URLError as e:
            return clasificar(url, None, None, f"Sin respuesta ({e.reason})")
        except Exception as e:                      # timeout, SSL, etc.
            return clasificar(url, None, None, f"Sin respuesta ({e})")
    return clasificar(url, None, None, "Sin respuesta")


# --------------------------------------------------------------- modo API
def clasificar_estado_api(estado: str | None,
                          sub_estados=None) -> tuple[str, str]:
    """Traduce el campo `status` de la API. Funcion pura."""
    if not estado:
        return SIN_VERIFICAR, "La API no devolvió el estado del aviso"
    estado = str(estado).lower()
    subs = [str(s).lower() for s in (sub_estados or [])]
    if estado == "active" and not any(
            s in ("deleted", "suspended", "expired") for s in subs):
        return VIGENTE, "La API lo reporta activo"
    detalle = f"{estado}" + (f" ({', '.join(subs)})" if subs else "")
    return CAIDO, f"La API lo reporta como {detalle}"


def revisar_con_api(api, urls: list[str],
                    avance=None) -> dict[str, tuple[str, str]]:
    """
    Consulta el estado real de cada aviso, 20 por llamada.
    Devuelve {url: (estado, motivo)}.
    """
    por_codigo: dict[str, list[str]] = {}
    salida: dict[str, tuple[str, str]] = {}
    for u in urls:
        cod = codigo_mco(u)
        if cod:
            por_codigo.setdefault(cod, []).append(u)
        else:
            salida[u] = (SIN_VERIFICAR, "El enlace guardado no tiene código de aviso")

    codigos = list(por_codigo)
    for i in range(0, len(codigos), 20):
        lote = codigos[i:i + 20]
        try:
            crudo = api.get("/items", ids=",".join(lote),
                            attributes="id,status,sub_status,permalink")
        except Exception as e:
            for c in lote:
                for u in por_codigo[c]:
                    salida[u] = (SIN_VERIFICAR, f"No se pudo consultar ({e})")
            if avance:
                avance(len(lote))
            continue

        vistos = set()
        for entrada in crudo if isinstance(crudo, list) else []:
            cuerpo = entrada.get("body") or {}
            cod = codigo_mco(str(cuerpo.get("id") or ""))
            if not cod or cod not in por_codigo:
                continue
            vistos.add(cod)
            if entrada.get("code") == 404:
                res = (CAIDO, "La API no encuentra el aviso")
            else:
                res = clasificar_estado_api(cuerpo.get("status"),
                                            cuerpo.get("sub_status"))
            for u in por_codigo[cod]:
                salida[u] = res
        for c in set(lote) - vistos:
            for u in por_codigo[c]:
                salida[u] = (CAIDO, "La API no encuentra el aviso")
        if avance:
            avance(len(lote))
    return salida


# ------------------------------------------------------------------- excel
def columna_de(hoja, *prefijos: str) -> int | None:
    """Numero de la primera columna cuyo encabezado empieza por algun prefijo."""
    import unicodedata

    def limpio(t):
        t = unicodedata.normalize("NFKD", str(t or ""))
        return "".join(c for c in t if not unicodedata.combining(c)).strip().lower()

    objetivo = [limpio(p) for p in prefijos]
    for celda in hoja[1]:
        if any(limpio(celda.value).startswith(p) for p in objetivo if p):
            return celda.column
    return None


def _abrir_api(modo: str):
    """Devuelve un cliente de la API si hay credenciales utilizables, o None."""
    if modo == "http":
        return None
    try:
        sys.path.insert(0, str(AQUI))
        from api_ml import ApiML, cargar_credenciales
        cid, sec = cargar_credenciales(AQUI / "credenciales.json")
        api = ApiML(cid, sec, verboso=False)
        api.token()                     # falla temprano si son invalidas
        return api
    except Exception:
        return None


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        description="Revisa si los enlaces del Excel siguen abriendo el aviso.")
    ap.add_argument("--archivo", type=Path, default=XLSX_POR_DEFECTO)
    ap.add_argument("--hilos", type=int, default=6,
                    help="revisiones en paralelo (por defecto 6)")
    ap.add_argument("--pausa", type=float, default=0.25,
                    help="segundos de espera entre revisiones de cada hilo")
    ap.add_argument("--solo-reporte", action="store_true",
                    help="no escribir nada en el Excel")
    ap.add_argument("--limite", type=int, default=0,
                    help="revisar solo los primeros N enlaces (para probar)")
    ap.add_argument("--modo", choices=["auto", "api", "http"], default="auto",
                    help="auto usa la API si hay credenciales; si no, HTTP")
    args = ap.parse_args(argv)

    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[!] Falta openpyxl.  Ejecuta:  pip install -r requirements.txt")
        return 2

    if not args.archivo.exists():
        print(f"[!] No encuentro {args.archivo}")
        return 2

    wb = load_workbook(args.archivo)
    if "Avisos" not in wb.sheetnames:
        print(f"[!] {args.archivo.name} no tiene una hoja llamada 'Avisos'")
        return 2
    ws = wb["Avisos"]

    col_link = columna_de(ws, "link", "enlace")
    if not col_link:
        print("[!] No encuentro la columna del enlace en la hoja 'Avisos'")
        return 2
    col_titulo = columna_de(ws, "titulo del aviso", "titulo")

    filas = []
    for f in range(2, ws.max_row + 1):
        url = ws.cell(f, col_link).value
        if isinstance(url, str) and url.strip().startswith("http"):
            filas.append((f, url.strip(),
                          ws.cell(f, col_titulo).value if col_titulo else ""))
    if args.limite:
        filas = filas[:args.limite]

    if not filas:
        print("[!] La hoja no tiene enlaces que revisar.")
        return 1

    print(f"\nRevisando {len(filas)} enlaces de {args.archivo.name}")
    print("Esto abre cada aviso para ver si todavía existe; "
          "puede tardar unos minutos.\n")

    resultados: dict[int, tuple[str, str]] = {}
    hechos = [0]

    def avanzar(cuantos=1):
        hechos[0] += cuantos
        n = hechos[0]
        if n % 10 < cuantos or n >= len(filas):
            print(f"  {min(n, len(filas)):>4}/{len(filas)}", flush=True)

    api = _abrir_api(args.modo)
    if api is not None:
        print("  Modo API: se consulta el estado real de cada aviso.\n")
        por_url = revisar_con_api(api, [u for _f, u, _t in filas], avanzar)
        for fila, url, _t in filas:
            resultados[fila] = por_url.get(
                url, (SIN_VERIFICAR, "no revisado"))
    else:
        if args.modo == "api":
            print("[!] Pediste modo API pero no hay credenciales válidas.")
            return 2
        print("  Modo HTTP: se abre cada enlace para ver dónde termina.")
        print("  (con credenciales de la API el resultado sería exacto)\n")

        def tarea(entrada):
            fila, url, _titulo = entrada
            time.sleep(args.pausa)
            estado, motivo = revisar(url)
            with _candado:
                resultados[fila] = (estado, motivo)
                avanzar()

        with ThreadPoolExecutor(max_workers=max(1, args.hilos)) as pool:
            list(pool.map(tarea, filas))

    # ------------------------------------------------------------- resumen
    conteo = {VIGENTE: 0, CAIDO: 0, SIN_VERIFICAR: 0}
    for estado, _ in resultados.values():
        conteo[estado] = conteo.get(estado, 0) + 1
    total = len(resultados) or 1

    print("\n" + "-" * 62)
    for estado in (VIGENTE, CAIDO, SIN_VERIFICAR):
        n = conteo.get(estado, 0)
        print(f"  {estado:<16} {n:>4}  ({100 * n / total:.0f} %)")
    print("-" * 62)

    # -------------------------------------------------------------- reporte
    with REPORTE.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["fila_excel", "titulo", "enlace", "estado", "motivo"])
        for fila, url, titulo in filas:
            estado, motivo = resultados.get(fila, (SIN_VERIFICAR, "no revisado"))
            w.writerow([fila, titulo, url, estado, motivo])
    print(f"\n  Detalle por aviso:  {REPORTE}")

    caidos = [(f, t, u) for f, u, t in filas
              if resultados.get(f, ("",))[0] == CAIDO]
    if caidos:
        print(f"\n  Primeros avisos caídos ({len(caidos)} en total):")
        for _fila, titulo, url in caidos[:8]:
            print(f"    · {str(titulo)[:52]:<52} {_corto(url, 46)}")

    # --------------------------------------------------------------- excel
    if args.solo_reporte:
        print("\n  El Excel no se modificó (--solo-reporte).\n")
        return 0

    respaldo = args.archivo.with_name(
        f"{args.archivo.stem}_respaldo_{datetime.now():%Y%m%d_%H%M}"
        f"{args.archivo.suffix}")
    shutil.copy2(args.archivo, respaldo)

    col_estado = columna_de(ws, "estado del enlace") or (ws.max_column + 1)
    col_fecha = columna_de(ws, "enlace verificado") or (ws.max_column + 2)
    ws.cell(1, col_estado).value = "Estado del enlace"
    ws.cell(1, col_fecha).value = "Enlace verificado el"
    hoy = datetime.now().strftime("%Y-%m-%d")
    for fila, _url, _t in filas:
        estado, _motivo = resultados.get(fila, (SIN_VERIFICAR, ""))
        ws.cell(fila, col_estado).value = estado
        ws.cell(fila, col_fecha).value = hoy
    wb.save(args.archivo)

    print(f"\n  Respaldo del Excel:  {respaldo.name}")
    print(f"  Excel marcado:       {args.archivo.name}")
    print("\n  Abre el tablero: ahora puedes filtrar por enlace vigente.\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
